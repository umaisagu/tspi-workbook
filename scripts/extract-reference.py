#!/usr/bin/env python3
"""Build lib/reference/workbook-v3.json from the lecturer's TMA3084 TSPi Workbook v3.0 (team master, member workbook).

The Reference page (/reference) shows the workbook's manual, standards and guidelines, and rows from
the worked SAMPLE workbook (fictional team "AUV"). Text is copied from the workbooks, not retyped.
Re-run after the lecturer publishes a new version:

  python3 scripts/extract-reference.py docs/workbook-v3/TSPi_Workbook_TMA3084_v3.0.xlsx \
      docs/workbook-v3/TSPi_Member_Workbook_TMA3084_v3.0.xlsx docs/workbook-v3/TSPi_Workbook_TMA3084_SAMPLE_v3.0.xlsx
"""
import datetime as dt
import json
import os
import sys

import openpyxl

OUT = os.path.join(os.path.dirname(__file__), "..", "lib", "reference", "workbook-v3.json")


def text(v):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def rows(ws, first, last, cols):
    """Rows first..last (inclusive) as lists of the given column letters; empty rows dropped."""
    out = []
    for r in range(first, last + 1):
        vals = [text(ws[f"{c}{r}"].value) for c in cols]
        if any(v is not None and not (isinstance(v, str) and v.startswith("=")) for v in vals):
            out.append([None if isinstance(v, str) and v.startswith("=") else v for v in vals])
    return out


def find_row(ws, col, startswith, after=1):
    for r in range(after, ws.max_row + 1):
        v = ws[f"{col}{r}"].value
        if isinstance(v, str) and v.strip().startswith(startswith):
            return r
    raise SystemExit(f"{ws.title}: no row starting with {startswith!r} in column {col}")


def manual(ws):
    """Sections 3-6 (steps), 7 (forms), 8 (worked examples), 9 (things that go wrong)."""
    stages = []
    for title in ("3.", "4.", "5.", "6."):
        r = find_row(ws, "A", title)
        steps = []
        for row in range(r + 2, ws.max_row + 1):
            n = ws[f"A{row}"].value
            if n is None or not str(n).strip().isdigit():
                break
            steps.append({"n": int(n), "step": text(ws[f"B{row}"].value), "what": text(ws[f"C{row}"].value),
                          "form": text(ws[f"D{row}"].value), "who": text(ws[f"E{row}"].value)})
        stages.append({"title": text(ws[f"A{r}"].value), "steps": steps})

    def table(title, cols):
        r = find_row(ws, "A", title)
        out = []
        for row in range(r + 2, ws.max_row + 1):
            if ws[f"B{row}"].value is None:
                break
            out.append([text(ws[f"{c}{row}"].value) for c in cols])
        return out

    intro = [{"label": text(ws[f"B{r}"].value), "text": text(ws[f"C{r}"].value)} for r in range(5, 9)]
    closing = text(ws[f"B{find_row(ws, 'A', '10.') + 1}"].value)
    return {
        "intro": intro,
        "stages": stages,
        "forms": [dict(zip(("form", "holds", "owner", "feeds"), x)) for x in table("7.", "BCDE")],
        "examples": [dict(zip(("situation", "record"), x)) for x in table("8.", "BC")],
        "troubleshooting": [dict(zip(("symptom", "fix"), x)) for x in table("9.", "BC")],
        "closing": closing,
    }


def blank(path):
    wb = openpyxl.load_workbook(path)
    readme, setup = wb["README"], wb["SETUP"]
    ph, dtp, pg, qg, roles = wb["PHASES"], wb["DEFECTTYPES"], wb["PLANGUIDE"], wb["QUALGUIDE"], wb["ROLES"]
    return {
        "title": text(readme["A1"].value),
        "subtitle": text(readme["A2"].value),
        "author": wb.properties.creator,
        "about": [{"label": text(readme[f"B{r}"].value), "text": text(readme[f"C{r}"].value)} for r in range(4, 11)
                  if readme[f"B{r}"].value],
        "milestones": [dict(zip(("id", "name", "week", "start", "due", "weight", "deliverables"),
                                [text(setup[f"{c}{r}"].value) for c in "ABCDEFH"])) for r in range(20, 24)],
        "thresholds": {
            "columns": ["M1", "M2", "M3", "M4"],
            "rows": [{"label": text(setup[f"A{r}"].value), "values": [text(setup[f"{c}{r}"].value) for c in "BCDE"],
                      "note": text(setup[f"H{r}"].value)} for r in range(27, 34)],
        },
        "manual": manual(wb["MANUAL"]),
        "roles": [{"role": text(roles[f"A{r}"].value), "owns": text(roles[f"B{r}"].value)} for r in range(8, 14)],
        "roleRule": text(roles["A27"].value),
        "phases": [dict(zip(("code", "name", "scope", "logTime", "injected", "removed", "development", "sizeBasis", "belongs"),
                            [text(ph[f"{c}{r}"].value) for c in "ABCDEFGHI"])) for r in range(5, 32)],
        "phaseNotes": [text(ph[f"A{r}"].value) for r in range(34, 39)],
        "defectTypes": [dict(zip(("code", "name", "covers", "example"), [text(dtp[f"{c}{r}"].value) for c in "ABCD"]))
                        for r in range(5, 20)],
        "defectTypeNote": text(dtp["A22"].value),
        "planGuide": {
            "notes": [{"label": text(pg[f"B{r}"].value), "text": text(pg[f"C{r}"].value)} for r in (5, 6, 7, 10)],
            "shares": [dict(zip(("group", "share", "phases"), [text(pg[f"{c}{r}"].value) for c in "BCD"])) for r in range(12, 21)],
            "rates": [dict(zip(("activity", "rate", "why"), [text(pg[f"{c}{r}"].value) for c in "BCD"])) for r in range(24, 32)],
            "aiWarning": text(pg["B34"].value),
        },
        "qualGuide": {
            "idea": text(qg["B5"].value),
            "measures": [dict(zip(("measure", "aim", "below"), [text(qg[f"{c}{r}"].value) for c in "BCD"])) for r in range(8, 22)],
            "fixOrder": text(qg["B24"].value),
        },
    }


def member(path, master):
    """The member workbook (one per student) and the master's TEAM sheet that consolidates them."""
    wb = openpyxl.load_workbook(path)
    readme = wb["README"]
    team = openpyxl.load_workbook(master)["TEAM"]
    return {
        "title": text(readme["A1"].value),
        "about": [{"label": text(readme[f"B{r}"].value), "text": text(readme[f"C{r}"].value)} for r in range(4, 9)
                  if readme[f"B{r}"].value],
        "routine": [text(readme[f"C{r}"].value) for r in range(11, 16) if readme[f"C{r}"].value],
        "sheets": [{"sheet": ws.title, "title": text(ws["A1"].value), "purpose": text(ws["A2"].value)} for ws in wb.worksheets],
        "consolidation": [{"label": text(team[f"A{r - 1}"].value), "text": text(team[f"A{r}"].value)} for r in (17, 20, 23)],
    }


def sample(path):
    wb = openpyxl.load_workbook(path)
    setup, si = wb["SETUP"], wb["FORM-SI"]
    return {
        "note": text(wb["README"]["B3"].value),
        "team": {k: text(setup[f"C{r}"].value) for k, r in
                 (("demonstrator", 8), ("labGroup", 9), ("name", 10), ("project", 11), ("client", 12), ("repository", 14), ("size", 16))},
        "members": [dict(zip(("matric", "name", "github", "primaryRole", "secondaryRole"), [text(si[f"{c}{r}"].value) for c in "BCFGH"]))
                    for r in range(9, 15) if si[f"C{r}"].value],
        "charter": [{"label": text(si[f"B{r}"].value), "text": text(si[f"C{r}"].value)} for r in range(18, 24)],
        "task": rows(wb["TASK"], 8, 15, "ABCDEFKLO"),
        "logt": rows(wb["LOGT"], 7, 14, "ABCEFGHIJLM"),
        "logd": rows(wb["LOGD"], 7, 12, "ABCDEFGHIJKLM"),
        "logtest": rows(wb["LOGTEST"], 7, 12, "ABCDEFGHIJKL"),
        "logai": rows(wb["LOGAI"], 7, 10, "ABCDEFGHIKLMNOP"),
        "logclient": rows(wb["LOGCLIENT"], 7, 10, "ABCDEFGHIJK"),
        "loggit": rows(wb["LOGGIT"], 7, 12, "ABCDEFGHIJKL"),
        "ccr": rows(wb["CCR"], 7, 10, "ABCDEFGHIJKLM"),
        "week": rows(wb["WEEK"], 8, 12, "ABCDEMN"),
    }


def main():
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    master, member_wb, sample_wb = sys.argv[1:]
    data = {"source": [os.path.basename(p) for p in sys.argv[1:]], **blank(master), "member": member(member_wb, master),
            "sample": sample(sample_wb)}
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
        f.write("\n")
    print(f"wrote {os.path.relpath(OUT)}")


if __name__ == "__main__":
    main()
